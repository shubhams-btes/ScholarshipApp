import openpyxl
from openpyxl.styles import Font,Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from io import BytesIO
from tests.models import Result
from decimal import Decimal

def build_results_workbook(schedule, results, include_score=False):
    """Build the results Excel for a schedule. `results` is an already-filtered queryset.
    include_score=True adds the Score column (internal/official use);
    default False omits it (for sharing with colleges)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Results_{schedule.college.name}"[:31]

    headers = ["ID", "Student Name", "Email", "College", "Contact Number"]
    if include_score:
        headers.append("Score")
        headers.append("Hall Ticket")
        headers.append("Roll No")
        headers.extend(["10th %", "12th %", "Graduation Degree", "Graduation %", "Masters Degree", "Masters %"])

    for col_index, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_index, value=header.upper())
        cell.font = Font(bold=True)

    for idx, result in enumerate(results, start=1):
        s = result.student
        row = [
            idx,
            s.name.upper() if s.name else "",
            s.email.upper() if s.email else "",
            result.exam_schedule.college.name.upper() if result.exam_schedule.college.name else "",
            s.mobile_number if s.mobile_number else "",
        ]
        if include_score:
            row.append(result.score)
            row.append(s.hall_ticket.upper() if s.hall_ticket else "")
            row.append(s.roll_no.upper() if s.roll_no else "")
            row.extend([
                s.tenth_percentage if s.tenth_percentage is not None else "",
                s.twelfth_percentage if s.twelfth_percentage is not None else "",
                s.get_graduation_degree_display() if s.graduation_degree else "",
                s.graduation_percentage if s.graduation_percentage is not None else "",
                s.get_masters_degree_display() if s.masters_degree else "",
                s.masters_percentage if s.masters_percentage is not None else "",
            ])

        for col_index, value in enumerate(row, start=1):
            ws.cell(row=idx + 1, column=col_index, value=value)

    return wb

def get_filtered_results(schedule, cutoff=None, top_n=None):
    results = Result.objects.filter(
        exam_schedule=schedule
    ).select_related("student", "exam_schedule__college").order_by("-score")

    if cutoff:
        try:
            results = results.filter(score__gte=int(cutoff))
        except (ValueError, TypeError):
            pass

    if top_n:
        try:
            results = results[:int(top_n)]
        except (ValueError, TypeError):
            pass

    return results




def talentquest_percentage(result):
    """Exam score as % out of 100."""
    if not result.total_questions:
        return Decimal("0")
    return round(Decimal(result.score) / Decimal(result.total_questions) * 100, 2)


def final_percentage(academic_pct, tq_pct, interview_pct):
    """Equal-thirds average of academics, TalentQuest, interview.
    Returns None if interview not done (final can't be computed yet)."""
    if interview_pct is None or academic_pct is None:
        return None
    return round((Decimal(academic_pct) + Decimal(tq_pct) + Decimal(interview_pct)) / 3, 2)


def scholarship_slab(final_pct):
    
    if final_pct is None:
        return (None, None)
    if final_pct >= 90:
        return ("A", 50)
    if final_pct >= 80:
        return ("B", 25)
    if final_pct >= 70:
        return ("C", 15)
    return ("—", 0)


def build_shortlist_workbook(schedule, results):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Shortlist_{schedule.college.name}"[:31]

    # ── Color definitions ──
    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")   # light green
    red_font = Font(color="CC0000")                                                          # red text
    header_font = Font(bold=True)
    banner_font = Font(bold=True, size=14)

        # ── Border: thin on all sides ──
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    headers = [
        "ID", "Hall Ticket", "Roll No", "Name", "Email", "Contact Number",
        "Talent Quest Test (Round-1) Result",
        "Tech. Interview Result",
        "Scholarship Offered Slab",
        "Scholarship %",
        "Stack Offered",
        "Final %",
    ]
    num_cols = len(headers)

    # ── Row 1: merged college banner (yellow) ──
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=num_cols)
    banner = ws.cell(row=1, column=1, value=f"{schedule.college.name} ")
    banner.font = banner_font
    banner.alignment = Alignment(horizontal="center")
    banner.fill = yellow_fill
    # Fill the merged cells too (merge only styles the top-left; fill the rest for full banner color)
    for col in range(1, num_cols + 1):
        ws.cell(row=1, column=col).fill = yellow_fill
        ws.cell(row=1, column=col).border = border

    # ── Row 2: column headers (yellow) ──
    for col_index, header in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=col_index, value=header.upper())
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.fill = yellow_fill
        cell.border = border

    # ── Build + sort rows (unchanged) ──
    built = []
    for result in results:
        s = result.student
        academic = s.academic_percentage
        tq_pct = talentquest_percentage(result)
        interview = getattr(result, "interview", None)
        interview_pct = interview.interview_percentage if interview else None
        final = final_percentage(academic, tq_pct, interview_pct)
        slab, scholarship = scholarship_slab(final)

        if interview is None:
            interview_result = "Pending"
        elif interview.status == "SELECTED":
            interview_result = "Selected"
        else:
            interview_result = "Not Forwarded"

        built.append({
            "student": s, "score": result.score,
            "interview_result": interview_result,
            "slab": slab if slab and slab != "—" else "NA",
            "scholarship": scholarship if scholarship else "",
            "stack": interview.get_stream_display() if interview and interview.stream else "",
            "final": final,
            "status": interview.status if interview else None,
        })

    SLAB_ORDER = {"A": 0, "B": 1, "C": 2, "NA": 3}
    built.sort(key=lambda r: (0 if r["status"] == "SELECTED" else 1, SLAB_ORDER.get(r["slab"], 3)))

    # ── Write data rows with conditional coloring ──
    for idx, r in enumerate(built, start=1):
        s = r["student"]
        row_values = [
            idx,
            s.hall_ticket.upper() if s.hall_ticket else "",
            s.roll_no.upper() if s.roll_no else "",
            s.name.upper() if s.name else "",
            s.email.upper() if s.email else "",
            s.mobile_number if s.mobile_number else "",
            r["score"],
            r["interview_result"],
            r["slab"],
            r["scholarship"],
            r["stack"],
            r["final"] if r["final"] is not None else "",
        ]

        excel_row = idx + 2   # banner + header offset

        # Determine row styling by interview status
        is_selected = r["status"] == "SELECTED"
        is_rejected = r["status"] == "REJECTED"

        for col_index, value in enumerate(row_values, start=1):
            cell = ws.cell(row=excel_row, column=col_index, value=value)
            if is_selected:
                cell.fill = green_fill          # Selected → green background
            if is_rejected:
                cell.font = red_font            # Not Forwarded → red font
            cell.border = border
            
    # ── Adjust column widths for better readability ──
    MIN_WIDTH = 10
    MAX_WIDTH = 22
    PADDING = 3
    HEADER_ROW_HEIGHT = 45   # taller, so wrapped header text fits
    DATA_ROW_HEIGHT = 20     # standard data row height

    # ── Column WIDTHS (per column) ──
    COLUMN_MIN = {"EMAIL": 28, "NAME": 22, "HALL TICKET": 14}
    for col_index, header in enumerate(headers, start=1):
        base = max(min(len(header) + PADDING, MAX_WIDTH), MIN_WIDTH)
        width = max(base, COLUMN_MIN.get(header.upper(), 0))
        ws.column_dimensions[get_column_letter(col_index)].width = width

    # ── Row HEIGHTS (per row, by row number) ──
    ws.row_dimensions[1].height = 30                    # banner row
    ws.row_dimensions[2].height = HEADER_ROW_HEIGHT     # header row (taller for wrapped text)
    # Data rows start at row 3; set each to a standard height
    total_data_rows = len(built)
    for r in range(3, 3 + total_data_rows):
        ws.row_dimensions[r].height = DATA_ROW_HEIGHT

    return wb